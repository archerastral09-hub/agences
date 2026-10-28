#!/usr/bin/env python3
"""Century 21 : lit le numéro RCS dans <site agence>/mentions_legales/
et l'écrit dans la colonne SIREN (feuilles « Agences » et « Contacts »).
Aucun repli : si pas de RCS trouvé, la case reste vide.

Usage : python recup_siren_century21.py data/agences_century21.xlsx -o sortie.xlsx
"""
import argparse, json, re, sys, time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl
import requests
from bs4 import BeautifulSoup

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/124 Safari/537.36",
      "Accept-Language": "fr-FR,fr;q=0.9"}
# « RCS : CRETEIL B 451 071 468 CRETEIL » -> 451071468 (9 chiffres exactement)
RCS_RE = re.compile(r"RCS[^\d]{0,60}?(?<!\d)((?:\d[ .]?){8}\d)(?!\d)", re.I)


def get(session, url):
    for essai in range(4):
        try:
            r = session.get(url, headers=UA, timeout=25)
            if r.status_code == 429:
                time.sleep(5 * (essai + 1)); continue
            return r
        except requests.RequestException:
            time.sleep(2 * (essai + 1))
    return None


def rcs_depuis_site(session, site):
    if not site:
        return None
    r = get(session, site.rstrip("/") + "/mentions_legales/")
    if not r or r.status_code != 200:
        return None
    soup = BeautifulSoup(r.text, "html.parser")
    bloc = soup.select_one(".c-the-legal-notices-agency-informations")
    texte = (bloc or soup).get_text("\n", strip=True).replace("\xa0", " ")
    texte = re.split(r"h[ée]bergeur", texte, flags=re.I)[0]  # ignore l'hébergeur
    m = RCS_RE.search(texte)
    return re.sub(r"\D", "", m.group(1)) if m else None


def traiter(t):
    ident, site = t
    return ident, rcs_depuis_site(requests.Session(), site)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out", default="agences_century21_siren.xlsx")
    ap.add_argument("--cache", default="siren_cache_century21.json")
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx)
    ws, wc = wb["Agences"], wb["Contacts"]
    h = {c.value: i for i, c in enumerate(ws[1], 1)}
    cache_p = Path(a.cache)
    cache = json.loads(cache_p.read_text()) if cache_p.exists() else {}

    todo = []
    for r in range(2, ws.max_row + 1):
        ident = ws.cell(r, h["ID agence"]).value
        if ident and ident not in cache and not ws.cell(r, h["SIREN"]).value:
            todo.append((ident, ws.cell(r, h["Site internet"]).value))
    if a.limit:
        todo = todo[:a.limit]
    print(f"{len(todo)} agences à traiter ({len(cache)} en cache)", flush=True)

    try:
        with ThreadPoolExecutor(a.workers) as ex:
            futs = [ex.submit(traiter, t) for t in todo]
            for n, f in enumerate(as_completed(futs), 1):
                ident, siren = f.result()
                cache[ident] = siren
                if n % 25 == 0:
                    cache_p.write_text(json.dumps(cache))
                    print(f"  {n}/{len(todo)}", flush=True)
    finally:
        cache_p.write_text(json.dumps(cache))

    par_id = {}
    for r in range(2, ws.max_row + 1):
        ident = ws.cell(r, h["ID agence"]).value
        if not ws.cell(r, h["SIREN"]).value and cache.get(ident):
            ws.cell(r, h["SIREN"]).value = cache[ident]
        par_id[ident] = ws.cell(r, h["SIREN"]).value
    hc = {c.value: i for i, c in enumerate(wc[1], 1)}
    for r in range(2, wc.max_row + 1):
        wc.cell(r, hc["SIREN"]).value = par_id.get(wc.cell(r, hc["ID agence"]).value)
    wb.save(a.out)

    ok = sum(1 for v in par_id.values() if v)
    print(f"RCS/SIREN trouvés : {ok}/{len(par_id)} -> {a.out}")
    vides = [i for i, v in par_id.items() if not v]
    if vides:
        print("Sans RCS :", ", ".join(vides[:50]), "…" if len(vides) > 50 else "")


if __name__ == "__main__":
    sys.exit(main())
