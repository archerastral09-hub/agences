#!/usr/bin/env python3
"""
Century 21 : complète agences_century21.xlsx
  1) visite la page /mentions_legales/ de chaque agence -> SIREN, SIRET, email
     (l'email est une IMAGE sur ces pages -> OCR tesseract)
  2) interroge l'API Recherche d'entreprises (api.gouv.fr / data.gouv.fr)
     avec le SIREN -> dirigeants, ajoutés dans la feuille « Contacts »

Usage : python -u scrape_century21.py [--input agences_century21.xlsx]
                                      [--output output/agences_century21_complete.xlsx]
                                      [--limit 20] [--workers 8]
"""
import argparse
import difflib
import io
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urljoin

import openpyxl
import pytesseract
import requests
from bs4 import BeautifulSoup
from PIL import Image, ImageOps
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
API_URL = "https://recherche-entreprises.api.gouv.fr/search"
API_SOURCE = "API Recherche d'entreprises (data.gouv.fr)"
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
DEFAULT_EMAIL_DOMAIN = "century21france.fr"

NEW_COLS = [
    "Raison sociale (API)",
    "Statut entreprise (API)",
    "Directeur publication (site)",
    "Email source",
    "Email OCR brut",
    "Statut scraping",
]


def log(msg):
    print(msg, flush=True)


def make_session():
    s = requests.Session()
    s.headers.update({"User-Agent": UA, "Accept-Language": "fr-FR,fr;q=0.9"})
    retry = Retry(total=3, backoff_factor=1.5,
                  status_forcelist=(429, 500, 502, 503, 504),
                  allowed_methods=("GET",))
    s.mount("https://", HTTPAdapter(max_retries=retry, pool_maxsize=32))
    s.mount("http://", HTTPAdapter(max_retries=retry, pool_maxsize=32))
    return s


# ----------------------------------------------------------------- SIREN ---
def luhn_ok(num: str) -> bool:
    total = 0
    for i, ch in enumerate(reversed(num)):
        d = int(ch)
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def split_siren_siret(digits: str):
    """Sur ces pages le champ « SIREN » contient souvent un SIRET (14 chiffres)."""
    if len(digits) == 14:
        return digits[:9], digits
    if len(digits) == 9:
        return digits, ""
    return "", ""


NUM_RE = r"((?:\d[\s.\u00a0]?){9,14})"
SIREN_RE = re.compile(r"SIRE[NT]\s*(?:n°|:|N°)?\s*" + NUM_RE, re.I)
RCS_RE = re.compile(r"RCS\s+[A-Za-zÀ-ÿ' -]{2,30}?\s*[:\-]?\s*" + NUM_RE, re.I)


def is_host_block(p) -> bool:
    """Exclut le bloc de l'hébergeur (Naxos...), qui a aussi un SIREN."""
    txt = p.get_text(" ", strip=True).lower()
    if "naxos" in txt or "hébergeur" in txt or "hebergeur" in txt:
        return True
    prev = p.find_previous(["h2", "h3", "h4"])
    return bool(prev and "hébergeur" in prev.get_text().lower())


def extract_siren(soup):
    for rx in (SIREN_RE, RCS_RE):
        for p in soup.find_all("p"):
            if is_host_block(p):
                continue
            m = rx.search(p.get_text(" ", strip=True))
            if m:
                digits = re.sub(r"\D", "", m.group(1))
                siren, siret = split_siren_siret(digits)
                if siren:
                    return siren, siret
    return "", ""


def extract_director(soup):
    for sp in soup.find_all(string=re.compile(r"Directeur de la publication", re.I)):
        p = sp.find_parent("p")
        if p:
            txt = p.get_text(" ", strip=True)
            txt = re.sub(r"(?i)directeur de la publication\s*:?", "", txt).strip()
            if txt:
                return txt
    return ""


# ----------------------------------------------------------------- EMAIL ---
def ocr_image(img_bytes: bytes) -> str:
    im = Image.open(io.BytesIO(img_bytes)).convert("RGBA")
    bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
    bg.alpha_composite(im)
    g = bg.convert("L")
    outs = []
    for scale, thr in ((3, None), (3, 170), (4, None)):
        x = g.resize((g.width * scale, g.height * scale), Image.LANCZOS)
        if thr:
            x = x.point(lambda v, t=thr: 255 if v > t else 0)
        x = ImageOps.expand(x, border=40, fill=255)
        txt = pytesseract.image_to_string(x, config="--psm 7").strip()
        if txt:
            outs.append(txt.replace(" ", ""))
    # le résultat qui contient un vrai email gagne, sinon le plus long
    for o in outs:
        if EMAIL_RE.fullmatch(o):
            return o
    return max(outs, key=len) if outs else ""


def expected_email_from_url(img_url: str) -> str:
    """.../imagesBien/202/590/contact.png -> ag590@century21france.fr"""
    m = re.search(r"/imagesBien/\d+/(\d+)/contact\.png", img_url or "")
    return f"ag{m.group(1)}@{DEFAULT_EMAIL_DOMAIN}" if m else ""


def resolve_email(raw_ocr: str, img_url: str):
    """
    L'OCR lit mal le « @ » et certains chiffres (ex. 'ag590facentury21france.fr').
    -> on vérifie l'OCR contre l'email attendu déduit de l'URL de l'image.
    Retourne (email, source).
    """
    if raw_ocr and EMAIL_RE.fullmatch(raw_ocr):
        return raw_ocr.lower(), "ocr"
    expected = expected_email_from_url(img_url)
    if expected and raw_ocr:
        ratio = difflib.SequenceMatcher(
            None, raw_ocr.lower(), expected.lower()).ratio()
        if ratio >= 0.75:
            return expected, "ocr+url"
    if expected:
        return expected, "déduit de l'URL (à vérifier)"
    return "", ""


# emails qui ne sont PAS ceux de l'agence (médiateur de la consommation, etc.)
EXCLUDED_DOMAINS = {"medimmoconso.fr"}


def is_excluded(mail: str) -> bool:
    return mail.lower().split("@")[-1] in EXCLUDED_DOMAINS


def find_email_block(soup):
    """Le <p> dont le libellé « Adresse mail » précède l'email de l'agence."""
    for p in soup.find_all("p"):
        label = p.find("strong")
        if label and "adresse mail" in label.get_text(" ", strip=True).lower():
            return p
    return None


def extract_email(soup, base_url, session):
    """Retourne (email, source, ocr_brut). Ne lit QUE le bloc « Adresse mail »."""
    p = find_email_block(soup)
    if p is None:
        return "", "", ""
    # 1) email en clair dans le bloc (mailto ou texte)
    for a in p.select('a[href^="mailto:"]'):
        mail = a["href"][7:].split("?")[0].strip()
        if EMAIL_RE.fullmatch(mail) and not is_excluded(mail):
            return mail.lower(), "texte", ""
    m = EMAIL_RE.search(p.get_text(" ", strip=True))
    if m and not is_excluded(m.group(0)):
        return m.group(0).lower(), "texte", ""
    # 2) email en image -> OCR
    img = p.find("img")
    if img and img.get("src"):
        img_url = urljoin(base_url, img["src"])
        r = session.get(img_url, timeout=30)
        r.raise_for_status()
        raw = ocr_image(r.content)
        email, source = resolve_email(raw, img_url)
        if email and is_excluded(email):
            return "", "", raw
        return email, source, raw
    return "", "", ""


# ------------------------------------------------------------------ PAGE ---
def parse_page(html: str, base_url: str, session):
    soup = BeautifulSoup(html, "html.parser")
    siren, siret = extract_siren(soup)
    email, esrc, raw = extract_email(soup, base_url, session)
    return {
        "siren": siren, "siret": siret,
        "siren_luhn": luhn_ok(siren) if siren else None,
        "email": email, "email_source": esrc, "email_ocr_brut": raw,
        "directeur_site": extract_director(soup),
    }


def scrape_agency(session, site: str):
    base = site.rstrip("/") + "/"
    url = urljoin(base, "mentions_legales/")
    r = session.get(url, timeout=30, allow_redirects=True)
    if r.status_code != 200:
        return {"statut": f"HTTP {r.status_code}"}
    data = parse_page(r.text, r.url, session)
    data["statut"] = "ok" if data["siren"] else "ok (SIREN introuvable)"
    return data


# ------------------------------------------------------------------- API ---
def fetch_dirigeants(session, siren: str):
    """API Recherche d'entreprises : /search?q=<siren>"""
    for attempt in range(5):
        r = session.get(API_URL, params={"q": siren, "page": 1, "per_page": 5},
                        timeout=30)
        if r.status_code == 429:
            time.sleep(2 * (attempt + 1))
            continue
        r.raise_for_status()
        for res in r.json().get("results", []):
            if res.get("siren") == siren:
                return res
        return None
    return None


def format_dirigeant(d):
    if d.get("type_dirigeant") == "personne morale":
        nom = d.get("denomination") or ""
        prenom = ""
        fonction = f"{d.get('qualite', '')} (personne morale"
        fonction += f", SIREN {d['siren']})" if d.get("siren") else ")"
    else:
        nom = (d.get("nom") or "").strip()
        prenom = (d.get("prenoms") or "").strip()
        fonction = d.get("qualite") or ""
    return nom, prenom, fonction


# ------------------------------------------------------------------ MAIN ---
def ensure_columns(ws, names):
    header = [c.value for c in ws[1]]
    idx = {}
    for n in names:
        if n in header:
            idx[n] = header.index(n) + 1
        else:
            ws.cell(row=1, column=ws.max_column + 1, value=n)
            header.append(n)
            idx[n] = len(header)
    return idx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default="agences_century21.xlsx")
    ap.add_argument("--output", default="output/agences_century21_complete.xlsx")
    ap.add_argument("--limit", type=int, default=0, help="0 = toutes les agences")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--api-delay", type=float, default=0.2)
    args = ap.parse_args()

    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
    wb = openpyxl.load_workbook(args.input)
    ws = wb["Agences"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    extra = ensure_columns(ws, NEW_COLS)

    rows = []
    for r in range(2, ws.max_row + 1):
        rows.append((r, ws.cell(r, col["ID agence"]).value,
                     (ws.cell(r, col["Site internet"]).value or "").strip()))
    if args.limit:
        rows = rows[:args.limit]
    log(f"{len(rows)} agences à traiter ({args.workers} threads)")

    session = make_session()

    # ---------- étape 1 : pages mentions légales ----------
    results = {}

    def work(item):
        r, ag_id, site = item
        if not site:
            return r, ag_id, {"statut": "pas de site"}
        try:
            return r, ag_id, scrape_agency(session, site)
        except Exception as e:  # une agence en erreur ne doit pas tout bloquer
            return r, ag_id, {"statut": f"erreur: {type(e).__name__}: {e}"[:200]}

    done = 0
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = [ex.submit(work, it) for it in rows]
        for f in as_completed(futs):
            r, ag_id, data = f.result()
            results[r] = (ag_id, data)
            done += 1
            log(f"[{done}/{len(rows)}] {ag_id} | {data.get('statut')} | "
                f"SIREN={data.get('siren', '')} | email={data.get('email', '')} "
                f"({data.get('email_source', '')})")

    for r, (ag_id, d) in results.items():
        if d.get("siren"):
            ws.cell(r, col["SIREN"], d["siren"])
        if d.get("siret"):
            ws.cell(r, col["SIRET"], d["siret"])
        if d.get("email"):
            ws.cell(r, col["Email général"], d["email"])
        ws.cell(r, extra["Email source"], d.get("email_source", ""))
        ws.cell(r, extra["Email OCR brut"], d.get("email_ocr_brut", ""))
        ws.cell(r, extra["Directeur publication (site)"], d.get("directeur_site", ""))
        ws.cell(r, extra["Statut scraping"], d.get("statut", ""))

    # ---------- étape 2 : API dirigeants ----------
    wc = wb["Contacts"]
    ch = [c.value for c in wc[1]]
    cc = {h: i + 1 for i, h in enumerate(ch)}
    existing = set()
    for r in range(2, wc.max_row + 1):
        existing.add((wc.cell(r, cc["ID agence"]).value,
                      (wc.cell(r, cc["Nom"]).value or "").strip().upper(),
                      (wc.cell(r, cc["Prénom"]).value or "").strip().upper()))

    api_session = make_session()
    sirens = {}
    for r, (ag_id, d) in results.items():
        if d.get("siren"):
            sirens[ag_id] = (r, d["siren"])
    log(f"API : {len(sirens)} SIREN à interroger")

    n_new = 0
    for i, (ag_id, (r, siren)) in enumerate(sirens.items(), 1):
        try:
            res = fetch_dirigeants(api_session, siren)
        except Exception as e:
            log(f"[API {i}/{len(sirens)}] {ag_id} {siren} erreur {e}")
            continue
        time.sleep(args.api_delay)  # limite API ~7 req/s
        if not res:
            log(f"[API {i}/{len(sirens)}] {ag_id} {siren} introuvable")
            continue
        ws.cell(r, extra["Raison sociale (API)"], res.get("nom_complet", ""))
        ws.cell(r, extra["Statut entreprise (API)"],
                {"A": "Active", "C": "Cessée"}.get(res.get("etat_administratif"),
                                                    res.get("etat_administratif", "")))
        dirs = res.get("dirigeants") or []
        for d in dirs:
            nom, prenom, fonction = format_dirigeant(d)
            key = (ag_id, nom.upper(), prenom.split(",")[0].strip().upper())
            if key in existing:
                continue
            existing.add(key)
            row = wc.max_row + 1
            wc.cell(row, cc["ID agence"], ag_id)
            wc.cell(row, cc["SIREN"], siren)
            wc.cell(row, cc["Nom"], nom)
            wc.cell(row, cc["Prénom"], prenom)
            wc.cell(row, cc["Fonction"], fonction)
            wc.cell(row, cc["Source"], API_SOURCE)
            n_new += 1
        log(f"[API {i}/{len(sirens)}] {ag_id} {siren} -> "
            f"{res.get('nom_complet', '')} | {len(dirs)} dirigeant(s)")

    # SIREN sur les contacts existants
    by_id = {ag: s for ag, (_, s) in sirens.items()}
    for r in range(2, wc.max_row + 1):
        ag = wc.cell(r, cc["ID agence"]).value
        if ag in by_id and not wc.cell(r, cc["SIREN"]).value:
            wc.cell(r, cc["SIREN"], by_id[ag])

    wb.save(args.output)
    log(f"Terminé : {len(sirens)} SIREN, {n_new} dirigeants ajoutés -> {args.output}")


if __name__ == "__main__":
    sys.exit(main())
  
