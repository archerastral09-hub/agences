#!/usr/bin/env python3
"""Récupère le SIREN de chaque agence Orpi et remplit les colonnes SIREN
des feuilles « Agences » et « Contacts » du classeur Excel.

Stratégie (par agence) :
  1. Page de l'agence (colonne « Source ») : bloc « Mentions légales »
     -> « Numéro RCS : 839400744 » (= SIREN). Validé par la clé de Luhn.
  2. Repli : API publique recherche-entreprises.api.gouv.fr
     (raison sociale + code postal / ville).

Usage : python recup_siren.py data/agences_orpi.xlsx -o sortie.xlsx
"""
import argparse, json, re, sys, time, unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl
import requests
from bs4 import BeautifulSoup

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/124 Safari/537.36",
      "Accept-Language": "fr-FR,fr;q=0.9"}
API = "https://recherche-entreprises.api.gouv.fr/search"


def luhn_ok(s: str) -> bool:
    if not re.fullmatch(r"\d{9}", s):
        return False
    tot = 0
    for i, c in enumerate(reversed(s)):
        d = int(c)
        if i % 2:
            d = d * 2 - 9 if d > 4 else d * 2
        tot += d
    return tot % 10 == 0


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9 ]", " ", s.lower()).strip()


def get(session, url, **kw):
    for essai in range(4):
        try:
            r = session.get(url, timeout=25, **kw)
            if r.status_code == 429:
                time.sleep(5 * (essai + 1)); continue
            return r
        except requests.RequestException:
            time.sleep(2 * (essai + 1))
    return None


def siren_depuis_page(session, url):
    """Renvoie (siren, raison_sociale) lus dans les mentions légales."""
    r = get(session, url, headers=UA)
    if not r or r.status_code != 200:
        return None, None
    soup = BeautifulSoup(r.text, "html.parser")
    bloc = soup.select_one("#dialog-agency-legal-mentions") or soup
    texte = bloc.get_text("\n", strip=True).replace("\xa0", " ")
    raison = None
    li = bloc.select_one("ul li")
    if li:
        raison = li.get_text(strip=True)
    # « Numéro RCS : 839400744 » ou « SIREN / SIRET : ... »
    for m in re.finditer(r"(?:RCS|SIREN|SIRET)[^\d\n]{0,40}((?:\d[ .]?){9,14})", texte, re.I):
        digits = re.sub(r"\D", "", m.group(1))
        if len(digits) >= 9 and luhn_ok(digits[:9]):
            return digits[:9], raison
    return None, raison


def siren_depuis_api(session, nom, adresse, raison=None):
    """Repli : recherche par raison sociale (ou nom) + code postal."""
    cp = re.search(r"\b(\d{5})\b", adresse or "")
    cp = cp.group(1) if cp else None
    requete = raison or re.sub(r"^orpi\s+", "", nom or "", flags=re.I)
    params = {"q": requete, "per_page": 10, "etat_administratif": "A"}
    if cp:
        params["code_postal"] = cp
    time.sleep(0.15)  # l'API limite à ~7 req/s
    r = get(session, API, params=params)
    if not r or r.status_code != 200:
        return None
    res = r.json().get("results", [])
    cible = norm(requete)
    for e in res:
        noms = [e.get("nom_complet"), e.get("nom_raison_sociale")]
        if any(norm(n or "") and (norm(n) in cible or cible in norm(n)) for n in noms):
            if luhn_ok(e.get("siren", "")):
                return e["siren"]
    return None


def traiter(row):
    ident, nom, adresse, url = row
    s = requests.Session()
    siren, raison, src = None, None, None
    if url:
        siren, raison = siren_depuis_page(s, url)
        src = "page agence" if siren else None
    if not siren:
        siren = siren_depuis_api(s, nom, adresse, raison)
        src = "api gouv" if siren else None
    return ident, siren, src


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--out", default="agences_orpi_siren.xlsx")
    ap.add_argument("--cache", default="siren_cache.json")
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--limit", type=int, default=0, help="test : n premières agences")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx)
    ws, wc = wb["Agences"], wb["Contacts"]
    h = {c.value: i for i, c in enumerate(ws[1], 1)}
    cache_p = Path(a.cache)
    cache = json.loads(cache_p.read_text()) if cache_p.exists() else {}

    todo = []
    for r in range(2, ws.max_row + 1):
        ident = ws.cell(r, h["ID agence"]).value
        if not ident or ident in cache or ws.cell(r, h["SIREN"]).value:
            continue
        todo.append((ident, ws.cell(r, h["Nom agence"]).value,
                     ws.cell(r, h["Adresse"]).value, ws.cell(r, h["Source"]).value))
    if a.limit:
        todo = todo[:a.limit]
    print(f"{len(todo)} agences à traiter ({len(cache)} déjà en cache)", flush=True)

    try:
        with ThreadPoolExecutor(a.workers) as ex:
            futs = [ex.submit(traiter, t) for t in todo]
            for n, f in enumerate(as_completed(futs), 1):
                ident, siren, src = f.result()
                cache[ident] = {"siren": siren, "source": src}
                if n % 25 == 0:
                    cache_p.write_text(json.dumps(cache, ensure_ascii=False))
                    print(f"  {n}/{len(todo)}", flush=True)
    finally:
        cache_p.write_text(json.dumps(cache, ensure_ascii=False))

    # écriture dans « Agences » puis « Contacts »
    par_id = {}
    for r in range(2, ws.max_row + 1):
        ident = ws.cell(r, h["ID agence"]).value
        if ident in cache and cache[ident]["siren"] and not ws.cell(r, h["SIREN"]).value:
            ws.cell(r, h["SIREN"]).value = cache[ident]["siren"]
        par_id[ident] = ws.cell(r, h["SIREN"]).value
    hc = {c.value: i for i, c in enumerate(wc[1], 1)}
    for r in range(2, wc.max_row + 1):
        wc.cell(r, hc["SIREN"]).value = par_id.get(wc.cell(r, hc["ID agence"]).value)
    wb.save(a.out)

    ok = sum(1 for v in par_id.values() if v)
    print(f"SIREN trouvés : {ok}/{len(par_id)} -> {a.out}")
    manquants = [i for i, v in par_id.items() if not v]
    if manquants:
        print("Sans SIREN :", ", ".join(manquants[:50]), "…" if len(manquants) > 50 else "")


if __name__ == "__main__":
    sys.exit(main())
