#!/usr/bin/env python3
"""
Côté Particuliers : complète SIREN et SIRET de agences_coteparticuliers.xlsx

Constat : la page /mentions-legales.html d'une agence affiche l'identité du
SIÈGE DU RÉSEAU (CPART FRANCE, SIREN 509062063), pas celle de l'agence.
On ne copie donc PAS ces numéros dans l'agence. Pour chaque agence :
  1) on lit la page mentions légales ; si elle affiche un autre SIREN que
     celui du réseau, on le retient (méthode « page »)
  2) sinon on cherche l'agence dans l'API Recherche d'entreprises
     (api.gouv.fr / data.gouv.fr) par nom + code postal + adresse, et on
     retient le meilleur candidat (méthode « API », avec un score)

Usage : python -u scrape_coteparticuliers.py [--input agences_coteparticuliers.xlsx]
            [--output output/agences_coteparticuliers_complete.xlsx]
            [--limit 10] [--workers 6] [--min-score 5]
"""
import argparse
import os
import re
import sys
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urljoin

import openpyxl
import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
API_URL = "https://recherche-entreprises.api.gouv.fr/search"
NETWORK_SIREN = {"509062063"}          # CPART FRANCE (siège du réseau)
NETWORK_NAMES = ("cpart france",)
BRAND_KEYS = ("coteparticuliers", "coteparticulier")  # texte normalisé

NEW_COLS = [
    "Méthode SIREN",
    "Score correspondance",
    "Raison sociale (API)",
    "Adresse établissement (API)",
    "Statut entreprise (API)",
    "Raison sociale (site)",
    "Statut scraping",
]


def log(msg):
    print(msg, flush=True)


def make_session():
    s = requests.Session()
    s.headers.update({"User-Agent": UA, "Accept-Language": "fr-FR,fr;q=0.9"})
    retry = Retry(total=3, backoff_factor=1.5,
                  status_forcelist=(429, 500, 502, 503, 504),
                  allowed_methods=("GET",))
    s.mount("https://", HTTPAdapter(max_retries=retry, pool_maxsize=16))
    return s


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", s.lower())


def norm_words(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


# ------------------------------------------------------------- PAGE SITE ---
def split_siren_siret(digits: str):
    if len(digits) == 14:
        return digits[:9], digits
    if len(digits) == 9:
        return digits, ""
    return "", ""


def parse_mentions(html: str) -> dict:
    """Lit la liste « Raison sociale / RCS / Siret » (hors bloc hébergement)."""
    soup = BeautifulSoup(html, "html.parser")
    out = {"raison_sociale": "", "siren": "", "siret": ""}
    for li in soup.select("li.mention-list__ele"):
        h3 = li.find_previous("h3")
        if h3 and "bergement" in h3.get_text().lower():
            continue  # bloc hébergeur (HUMAN IMMOBILIER) : pas l'agence
        txt = li.get_text(" ", strip=True)
        low = txt.lower()
        if low.startswith("raison sociale"):
            out["raison_sociale"] = txt.split(":", 1)[-1].strip()
        elif low.startswith("siret"):
            d = re.sub(r"\D", "", txt.split(":", 1)[-1])
            s9, s14 = split_siren_siret(d)
            out["siret"] = s14 or out["siret"]
            out["siren"] = out["siren"] or s9
        elif low.startswith("rcs") or low.startswith("siren"):
            d = re.sub(r"\D", "", txt.split(":", 1)[-1])
            s9, _ = split_siren_siret(d[-9:] if len(d) > 9 else d)
            out["siren"] = out["siren"] or s9
    return out


def is_network(info: dict) -> bool:
    if info["siren"] in NETWORK_SIREN:
        return True
    rs = (info["raison_sociale"] or "").lower()
    return any(n in rs for n in NETWORK_NAMES)


def scrape_page(session, site: str) -> dict:
    url = urljoin(site.rstrip("/") + "/", "mentions-legales.html")
    r = session.get(url, timeout=30)
    if r.status_code != 200:
        return {"statut": f"HTTP {r.status_code}"}
    info = parse_mentions(r.text)
    info["statut"] = "ok"
    return info


# ------------------------------------------------------------------- API ---
def api_search(session, params):
    for attempt in range(5):
        r = session.get(API_URL, params=params, timeout=30)
        if r.status_code == 429:
            time.sleep(2 * (attempt + 1))
            continue
        r.raise_for_status()
        return r.json().get("results", [])
    return []


def split_address(addr: str):
    """'15 boulevard de la République, 47000 AGEN' -> (voie, cp, commune)"""
    m = re.search(r"(\d{5})\s+([^,]+)$", addr or "")
    cp = m.group(1) if m else ""
    commune = m.group(2).strip() if m else ""
    voie = (addr or "")[: m.start()].strip(" ,") if m else (addr or "")
    return voie, cp, commune


def street_key(voie: str):
    """numéro + premiers mots significatifs de la rue, normalisés."""
    m = re.match(r"\s*(\d+)", voie or "")
    num = m.group(1) if m else ""
    words = [w for w in norm_words(voie).split()
             if w not in ("rue", "avenue", "av", "boulevard", "bd", "place",
                          "de", "du", "la", "le", "des", "bis", "ter", "d")
             and not w.isdigit()]
    return num, words[:2]


def score_candidate(res, etab, nom_agence, voie, cp, commune):
    """Score de correspondance entre l'agence du fichier et un établissement."""
    if res.get("siren") in NETWORK_SIREN:
        return -99
    sc = 0
    if etab.get("code_postal") == cp:
        sc += 3
    if norm(etab.get("libelle_commune", "")) == norm(commune):
        sc += 1
    blob = norm(" ".join([
        res.get("nom_complet", "") or "",
        etab.get("nom_commercial", "") or "",
        " ".join(etab.get("liste_enseignes") or []),
    ]))
    if any(k in blob for k in BRAND_KEYS):
        sc += 3
    num, words = street_key(voie)
    adr = norm_words(etab.get("adresse", ""))
    if num and re.search(rf"\b{num}\b", adr):
        sc += 1
    if words and all(w in adr for w in words):
        sc += 2
    if str(res.get("activite_principale", "")).startswith("68.31"):
        sc += 1
    if etab.get("etat_administratif") == "A":
        sc += 1
    return sc


def find_in_api(session, nom_agence, adresse, api_delay):
    voie, cp, commune = split_address(adresse)
    queries = [
        {"q": nom_agence, "code_postal": cp},
        {"q": f"cote particuliers {commune}", "code_postal": cp},
        {"q": f"{voie} {commune}", "code_postal": cp,
         "activite_principale": "68.31Z"},
    ]
    best, best_sc = None, -99
    for p in queries:
        p = {**p, "per_page": 10, "page": 1}
        try:
            results = api_search(session, p)
        except Exception:
            results = []
        time.sleep(api_delay)
        for res in results:
            etabs = res.get("matching_etablissements") or []
            if res.get("siege"):
                etabs = etabs + [res["siege"]]
            for etab in etabs:
                sc = score_candidate(res, etab, nom_agence, voie, cp, commune)
                if sc > best_sc:
                    best, best_sc = (res, etab), sc
        if best_sc >= 9:   # correspondance quasi certaine : inutile d'insister
            break
    return best, best_sc


# ------------------------------------------------------------------ MAIN ---
def ensure_columns(ws, names):
    header = [c.value for c in ws[1]]
    idx = {}
    for n in names:
        if n in header:
            idx[n] = header.index(n) + 1
        else:
            ws.cell(row=1, column=ws.max_column + 1, value=n)
            header.append(n)
            idx[n] = len(header)
    return idx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", default="agences_coteparticuliers.xlsx")
    ap.add_argument("--output",
                    default="output/agences_coteparticuliers_complete.xlsx")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--api-delay", type=float, default=0.2)
    ap.add_argument("--min-score", type=int, default=5,
                    help="score minimum pour accepter un candidat API")
    args = ap.parse_args()

    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
    wb = openpyxl.load_workbook(args.input)
    ws = wb["Agences"]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}
    extra = ensure_columns(ws, NEW_COLS)

    rows = []
    for r in range(2, ws.max_row + 1):
        rows.append((r, ws.cell(r, col["ID agence"]).value,
                     ws.cell(r, col["Nom agence"]).value or "",
                     ws.cell(r, col["Adresse"]).value or "",
                     (ws.cell(r, col["Site internet"]).value or "").strip()))
    if args.limit:
        rows = rows[:args.limit]
    log(f"{len(rows)} agences à traiter")

    session = make_session()

    # ---------- étape 1 : pages mentions légales ----------
    pages = {}

    def work(item):
        r, ag_id, nom, adr, site = item
        if not site:
            return r, {"statut": "pas de site"}
        try:
            return r, scrape_page(session, site)
        except Exception as e:
            return r, {"statut": f"erreur: {type(e).__name__}: {e}"[:200]}

    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = [ex.submit(work, it) for it in rows]
        for i, f in enumerate(as_completed(futs), 1):
            r, info = f.result()
            pages[r] = info
            log(f"[page {i}/{len(rows)}] ligne {r} | {info.get('statut')} | "
                f"{info.get('raison_sociale', '')} | siren={info.get('siren', '')}")

    # ---------- étape 2 : décision par agence (+ API si besoin) ----------
    api_session = make_session()
    n_page = n_api = n_none = 0
    contact_sirens = {}
    for i, (r, ag_id, nom, adr, site) in enumerate(rows, 1):
        info = pages.get(r, {})
        ws.cell(r, extra["Raison sociale (site)"], info.get("raison_sociale", ""))
        ws.cell(r, extra["Statut scraping"], info.get("statut", ""))
        siren = siret = ""
        methode = ""
        if info.get("siren") and not is_network(info):
            siren, siret, methode = info["siren"], info.get("siret", ""), "page"
            n_page += 1
        else:
            best, sc = find_in_api(api_session, nom, adr, args.api_delay)
            if best and sc >= args.min_score:
                res, etab = best
                siren = res.get("siren", "")
                siret = etab.get("siret", "")
                methode = "API nom+adresse" + ("" if sc >= 9 else " (à vérifier)")
                ws.cell(r, extra["Score correspondance"], sc)
                ws.cell(r, extra["Raison sociale (API)"], res.get("nom_complet", ""))
                ws.cell(r, extra["Adresse établissement (API)"],
                        etab.get("adresse", ""))
                ws.cell(r, extra["Statut entreprise (API)"],
                        {"A": "Active", "C": "Cessée"}.get(
                            res.get("etat_administratif"),
                            res.get("etat_administratif", "")))
                n_api += 1
            else:
                methode = "non trouvé"
                if best:
                    ws.cell(r, extra["Score correspondance"], sc)
                n_none += 1
        ws.cell(r, extra["Méthode SIREN"], methode)
        if siren:
            ws.cell(r, col["SIREN"], siren)
            contact_sirens[ag_id] = siren
        if siret:
            ws.cell(r, col["SIRET"], siret)
        log(f"[{i}/{len(rows)}] {ag_id} {nom} -> SIREN={siren} SIRET={siret} "
            f"({methode})")

    # SIREN sur l'onglet Contacts
    if "Contacts" in wb.sheetnames:
        wc = wb["Contacts"]
        cc = {c.value: i + 1 for i, c in enumerate(wc[1])}
        for r in range(2, wc.max_row + 1):
            ag = wc.cell(r, cc["ID agence"]).value
            if ag in contact_sirens and not wc.cell(r, cc["SIREN"]).value:
                wc.cell(r, cc["SIREN"], contact_sirens[ag])

    wb.save(args.output)
    log(f"Terminé : page={n_page} API={n_api} non trouvés={n_none} "
        f"-> {args.output}")


if __name__ == "__main__":
    sys.exit(main())
